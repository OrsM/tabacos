import pytest
import openpyxl
from tabacos.brandsales import parse_brand_workbook
from tabacos.errors import UnsupportedLayout

MONTHS = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]


def workbook(path, year=2025, unit="Cajetillas (20 uds.)", brands=None, offset=0, totals=True):
    """A sheet laid out like the Comisionado's: titles, a header row, one row per brand, a Totales row."""
    brands = brands or {"ALLURE": [100 + m for m in range(12)], "CAMEL": [1000 + 10 * m for m in range(12)]}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1 + offset, 1 + offset, "CIGARRILLOS (Península e Illes Balears)")
    ws.cell(1 + offset, 8 + offset, f"AÑO {year}")
    ws.cell(2 + offset, 1 + offset, f"CUOTA DE VENTAS POR MARCAS ({unit})")
    for i, label in enumerate(["Marcas"] + MONTHS + ["Totales"]):
        ws.cell(4 + offset, 1 + offset + i, label)
    for r, (brand, vals) in enumerate(brands.items()):
        ws.cell(5 + offset + r, 1 + offset, brand)
        for i, v in enumerate(vals):
            ws.cell(5 + offset + r, 2 + offset + i, v)
        ws.cell(5 + offset + r, 14 + offset, sum(v for v in vals if v))
    if totals:
        row = 5 + offset + len(brands) + 1
        ws.cell(row, 1 + offset, "Totales")
        for i in range(12):
            ws.cell(row, 2 + offset + i, sum((v[i] or 0) for v in brands.values()))
        ws.cell(row, 14 + offset, sum(sum(x for x in v if x) for v in brands.values()))
    wb.save(path)


def test_reads_year_unit_and_a_value_per_brand_and_month(tmp_path):
    workbook(tmp_path / "a.xlsx")
    unit, rows = parse_brand_workbook(tmp_path / "a.xlsx")
    assert unit == "packs"
    got = {(r["month"], r["brand"]): r["value"] for r in rows}
    assert got[("2025-01", "ALLURE")] == 100 and got[("2025-12", "CAMEL")] == 1110 and len(rows) == 24


def test_euros_file_and_a_table_that_starts_in_column_b(tmp_path):
    workbook(tmp_path / "b.xlsx", year=2016, unit="Euros", offset=1)
    unit, rows = parse_brand_workbook(tmp_path / "b.xlsx")
    assert unit == "euros" and rows[0]["month"] == "2016-01"


def test_months_with_no_value_yet_are_left_out(tmp_path):
    workbook(tmp_path / "c.xlsx", year=2026, brands={"ALLURE": [5, 6, 7] + [None] * 9})
    _, rows = parse_brand_workbook(tmp_path / "c.xlsx")
    assert [r["month"] for r in rows] == ["2026-01", "2026-02", "2026-03"]


def test_brand_row_that_does_not_add_up_to_its_total_raises(tmp_path):
    workbook(tmp_path / "d.xlsx")
    wb = openpyxl.load_workbook(tmp_path / "d.xlsx")
    wb.active.cell(5, 14, 99999)                      # ALLURE's printed total
    wb.save(tmp_path / "d.xlsx")
    with pytest.raises(UnsupportedLayout):
        parse_brand_workbook(tmp_path / "d.xlsx")


def test_monthly_sums_must_match_the_totales_row(tmp_path):
    workbook(tmp_path / "e.xlsx")
    wb = openpyxl.load_workbook(tmp_path / "e.xlsx")
    wb.active.cell(8, 2, 5)                           # January of the Totales row
    wb.save(tmp_path / "e.xlsx")
    with pytest.raises(UnsupportedLayout):
        parse_brand_workbook(tmp_path / "e.xlsx")


def test_a_sheet_without_a_brand_table_raises(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "provincias"
    wb.save(tmp_path / "f.xlsx")
    with pytest.raises(UnsupportedLayout):
        parse_brand_workbook(tmp_path / "f.xlsx")


def test_a_second_table_below_the_first_totales_row_is_ignored(tmp_path):
    workbook(tmp_path / "g.xlsx")
    wb = openpyxl.load_workbook(tmp_path / "g.xlsx")
    ws = wb.active
    ws.cell(10, 1, "ALLURE")                          # same brand name, another territory, below the Totales row (row 8)
    for i in range(12):
        ws.cell(10, 2 + i, 7)
    wb.save(tmp_path / "g.xlsx")
    _, rows = parse_brand_workbook(tmp_path / "g.xlsx")
    assert {(r["month"], r["brand"]): r["value"] for r in rows}[("2025-01", "ALLURE")] == 100 and len(rows) == 24


def test_months_where_every_brand_is_zero_are_left_out_like_empty_ones(tmp_path):
    workbook(tmp_path / "h.xlsx", year=2026, brands={"ALLURE": [5, 6, 7] + [0] * 9, "CAMEL": [1, 2, 3] + [0] * 9})
    _, rows = parse_brand_workbook(tmp_path / "h.xlsx")
    assert sorted({r["month"] for r in rows}) == ["2026-01", "2026-02", "2026-03"]
