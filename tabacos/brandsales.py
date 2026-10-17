"""Monthly cigarette sales by brand from the Comisionado's Excel files (Estadisticas_Ventas), in packs or euros."""
import re
from pathlib import Path

import openpyxl

from .errors import UnsupportedLayout

MONTHS = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
_YEAR = re.compile(r"A[ÑN]O\s*(\d{4})", re.I)


def _number(v):
    if isinstance(v, (int, float)):
        return float(v)
    return None


def _close(a: float, b: float) -> bool:
    return abs(a - b) <= max(1.0, 0.001 * abs(b))


def parse_brand_workbook(path: Path) -> tuple[str, list[dict]]:
    """(unit, rows) from one workbook: unit is "packs" (cajetillas of 20) or "euros"; a row is {month, brand, value}.

    Every brand's months must add up to its printed total, and the brands to the Totales row when there is one.
    Months with no value yet (the current year) are left out. Raises UnsupportedLayout for any other sheet.
    """
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    rows = [list(r) for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)]
    wb.close()
    head = " ".join(str(c) for r in rows[:6] for c in r if c is not None)
    if not (year := _YEAR.search(head)):
        raise UnsupportedLayout("no 'AÑO yyyy' in the first rows")
    unit = "euros" if re.search(r"\(\s*Euros", head, re.I) else "packs" if re.search(r"Cajetilla", head, re.I) else None
    if unit is None:
        raise UnsupportedLayout("cannot tell packs from euros")
    pos = next(((i, j) for i, r in enumerate(rows[:12]) for j, c in enumerate(r) if str(c).strip().lower() == "marcas"), None)
    if pos is None:
        raise UnsupportedLayout("no 'Marcas' header")
    h, c0 = pos
    labels = [str(c).strip().lower() if c is not None else "" for c in rows[h][c0 + 1:c0 + 14]]
    if labels[:12] != MONTHS or labels[12] != "totales":
        raise UnsupportedLayout(f"unexpected month columns {labels}")
    brands: dict[str, list] = {}
    totals_row = None
    for r in rows[h + 1:]:
        name = str(r[c0]).strip() if len(r) > c0 and r[c0] is not None else ""
        cells = [(_number(r[c0 + 1 + i]) if len(r) > c0 + 1 + i else None) for i in range(13)]
        if not name or all(v is None for v in cells[:12]):
            continue
        if name.lower() == "totales":                      # the first table ends here; a second one (other territory) may follow
            totals_row = cells
            break
        acc = brands.setdefault(name, [None] * 13)
        for i, v in enumerate(cells):
            if v is not None:
                acc[i] = (acc[i] or 0.0) + v                      # a brand listed twice is summed
    for name, vals in brands.items():
        printed = vals[12]
        if printed is not None and not _close(sum(v for v in vals[:12] if v is not None), printed):
            raise UnsupportedLayout(f"{name}: months {sum(v for v in vals[:12] if v is not None):.0f} do not add up to total {printed:.0f}")
    if totals_row:
        for i in range(12):
            summed = sum(vals[i] for vals in brands.values() if vals[i] is not None)
            if totals_row[i] is not None and not _close(summed, totals_row[i]):
                raise UnsupportedLayout(f"{MONTHS[i]}: brands add up to {summed:.0f}, Totales row says {totals_row[i]:.0f}")
    live = [i for i in range(12) if any(vals[i] for vals in brands.values())]          # a month nobody sold in has not happened yet
    out = [{"month": f"{year[1]}-{i + 1:02d}", "brand": name, "value": vals[i]}
           for name, vals in brands.items() for i in live if vals[i] is not None]
    if not out:
        raise UnsupportedLayout("no brand rows")
    return unit, out
